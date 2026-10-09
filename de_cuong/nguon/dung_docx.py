# -*- coding: utf-8 -*-
"""Dựng đề cương (.docx) từ khung đề cương cũ.

Chạy:  python3 dung_docx.py          -> bản chữ đỏ (phần thêm mới hoặc sửa tô đỏ, có dòng ghi chú)
       python3 dung_docx.py --sach  -> bản sạch để nộp (toàn bộ chữ đen, bỏ dòng ghi chú)
"""
import copy
import re
import sys
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.shared import Cm
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from lxml import etree

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
SACH = "--sach" in sys.argv
from noi_dung import (GHI_CHU, OLD_NUM, REFS, SEC_KE_HOACH, SEC_NOI_DUNG,  # noqa: E402
                      TITLE_EN, TITLE_VI, K, N)

BASE = str(HERE / "khung_de_cuong.docx")
OUT = str(HERE.parent / ("23521228_BuiLeHuyPhuoc_DeCuongKLTN.docx" if SACH else "23521228_BuiLeHuyPhuoc_DeCuongKLTN_chu_do.docx"))
RED = "FF0000"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
REF_RE = re.compile(r"\{ref:(\w+)\}")


def q(tag):
    return "{%s}%s" % (W, tag)


doc = docx.Document(BASE)
tbl = doc.tables[1]
trs = [r._tr for r in tbl.rows]
assert len(trs) == 98, len(trs)


def ppr_of(i, j):
    p = trs[i].findall(".//" + q("p"))[j]
    return copy.deepcopy(p.find(q("pPr")))


PPR = {
    "text": ppr_of(4, 0),
    "label": ppr_of(5, 0),
    "dash": ppr_of(5, 1),
    "plus": ppr_of(57, 2),
    "month": ppr_of(57, 1),
    "label_keep": ppr_of(57, 0),
    "ref_label": ppr_of(91, 0),
    "ref": ppr_of(91, 1),
}
PPR["sub"] = copy.deepcopy(PPR["label"])
lt = copy.deepcopy(PPR["label"])
lt.remove(lt.find(q("keepNext")))
PPR["label_text"] = lt
cap = copy.deepcopy(PPR["text"])
cap.find(q("jc")).set(q("val"), "center")
ind = cap.find(q("ind"))
if ind is not None:
    cap.remove(ind)
PPR["caption"] = cap
img = copy.deepcopy(cap)
kn = etree.Element(q("keepNext"))
img.insert(0, kn)
PPR["image"] = img
# Mọi gạch đầu dòng "-" dùng chung numId 7, mọi dấu "+" dùng chung numId 8 (như khung cũ).
PPR["dash"].find(q("numPr")).find(q("numId")).set(q("val"), "7")
PPR["plus"].find(q("numPr")).find(q("numId")).set(q("val"), "8")

# ---------------- Đánh số trích dẫn theo thứ tự xuất hiện ----------------
order = []
for section in (SEC_NOI_DUNG, SEC_KE_HOACH):
    for row in section:
        for par in row:
            kind, payload = par
            if kind == "image":
                continue
            for run in payload:
                for m in REF_RE.finditer(run[1]):
                    if m.group(1) not in order:
                        order.append(m.group(1))
NUM = {k: i + 1 for i, k in enumerate(order)}
missing = [k for k in order if k not in REFS]
unused = [k for k in REFS if k not in NUM]
assert not missing, missing
assert not unused, unused


def expand(run):
    flag, text, b, i = run
    red = flag == "N"
    out, pos = [], 0
    for m in REF_RE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], b, i, red))
        key = m.group(1)
        n = NUM[key]
        out.append((f"[{n}]", b, i, red or OLD_NUM.get(key) != n))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], b, i, red))
    return out


def make_run(text, bold=False, italic=False, red=False, size=26):
    r = etree.Element(q("r"))
    rpr = etree.SubElement(r, q("rPr"))
    f = etree.SubElement(rpr, q("rFonts"))
    for a in ("ascii", "cs", "eastAsia", "hAnsi"):
        f.set(q(a), "Times New Roman")
    if bold:
        etree.SubElement(rpr, q("b"))
        etree.SubElement(rpr, q("bCs"))
    if italic:
        etree.SubElement(rpr, q("i"))
        etree.SubElement(rpr, q("iCs"))
    if red and not SACH:
        etree.SubElement(rpr, q("color")).set(q("val"), RED)
    etree.SubElement(rpr, q("sz")).set(q("val"), str(size))
    etree.SubElement(rpr, q("szCs")).set(q("val"), str(size))
    t = etree.SubElement(r, q("t"))
    t.text = text
    t.set(XML_SPACE, "preserve")
    return r


def make_par(kind, runs):
    p = etree.Element(q("p"))
    p.append(copy.deepcopy(PPR[kind]))
    for run in runs:
        for t, b, i, red in expand(run):
            p.append(make_run(t, b, i, red))
    return p


def set_runs(p, runs):
    """Thay toàn bộ run của một đoạn có sẵn, giữ nguyên pPr."""
    for r in p.findall(q("r")):
        p.remove(r)
    for run in runs:
        for t, b, i, red in expand(run):
            p.append(make_run(t, b, i, red))


ROW_TEMPLATE = copy.deepcopy(trs[5])


def make_row(top, bottom, can_split):
    tr = copy.deepcopy(ROW_TEMPLATE)
    tc = tr.find(q("tc"))
    for p in tc.findall(q("p")):
        tc.remove(p)
    tr.find(q("trPr")).find(q("cantSplit")).set(q("val"), "0" if can_split else "1")
    tcpr = tc.find(q("tcPr"))
    borders = tcpr.find(q("tcBorders"))
    for side, on in (("top", top), ("bottom", bottom)):
        e = borders.find(q(side))
        e.set(q("val"), "single" if on else "nil")
        e.set(q("sz"), "4" if on else "0")
    mar = tcpr.find(q("tcMar"))
    mar.find(q("top")).set(q("w"), "55.0" if top else "0.0")
    mar.find(q("bottom")).set(q("w"), "55.0" if bottom else "0.0")
    return tr


sig_tr = trs[97]
pending_images = []


def add_section(rows):
    n = len(rows)
    for idx, row in enumerate(rows):
        size = sum(len(r[1]) for par in row if par[0] != "image" for r in par[1])
        tr = make_row(top=(idx == 0), bottom=(idx == n - 1), can_split=size > 700)
        tc = tr.find(q("tc"))
        for kind, payload in row:
            if kind == "image":
                p = OxmlElement("w:p")
                p.append(copy.deepcopy(PPR["image"]))
                tc.append(p)
                pending_images.append((p, tc, payload))
            else:
                tc.append(make_par(kind, payload))
        # Tiêu đề mục luôn nằm cùng hàng với ý đầu tiên của mục, nên không cần "giữ với đoạn sau";
        # thuộc tính này làm các hàng bị xích vào nhau và để lại khoảng trắng lớn cuối trang.
        # Hàng dài được phép ngắt trang nên bỏ cả "giữ các dòng cùng nhau".
        tags = ("keepNext", "keepLines") if size > 700 else ("keepNext",)
        for ppr in tc.iter(q("pPr")):
            for tag in tags:
                e = ppr.find(q(tag))
                if e is not None:
                    ppr.remove(e)
        sig_tr.addprevious(tr)


# ---------------- Ghi chú đầu trang ----------------
note = [p for p in doc.paragraphs if p.text.startswith("Bản rà soát")]
assert len(note) == 1
if SACH:
    note[0]._p.getparent().remove(note[0]._p)
else:
    set_runs(note[0]._p, [N(GHI_CHU, i=True)])

# ---------------- Tên đề tài ----------------
title_ps = trs[0].findall(q("tc") + "/" + q("p"))
set_runs(title_ps[1], [K("Tiếng Việt: ", b=True), N(TITLE_VI, b=True)])
set_runs(title_ps[2], [K("Tiếng Anh: ", b=True), N(TITLE_EN, b=True)])

# ---------------- Xóa nội dung cũ (hàng 4..96), dựng nội dung mới ----------------
for tr in trs[4:97]:
    tbl._tbl.remove(tr)

add_section(SEC_NOI_DUNG)

ref_rows = []
for idx, key in enumerate(order):
    n = NUM[key]
    old = key in OLD_NUM
    runs = [N(f"[{n}] ") if (not old or OLD_NUM[key] != n) else K(f"[{n}] ")]
    for text, italic in REFS[key]:
        runs.append(K(text, i=italic) if old else N(text, i=italic))
    par = ("ref", runs)
    if idx == 0:
        ref_rows.append([("ref_label", [K("Tài liệu tham khảo", b=True)]), par])
    else:
        ref_rows.append([par])
add_section(SEC_KE_HOACH + ref_rows)

# ---------------- Hình ----------------
for p, tc, path in pending_images:
    para = Paragraph(p, _Cell(tc, tbl))
    para.add_run().add_picture(path, width=Cm(15.0))

# ---------------- Ngày ký ----------------
sig_ps = sig_tr.findall(q("tc"))[1].findall(q("p"))
set_runs(sig_ps[0], [K("TP. HCM, ngày ", b=True), N("…..", b=True), K(" tháng ", b=True),
                     N("10", b=True), K(" năm 2026", b=True)])

# ---------------- Làm sạch thuộc tính số thực do khung cũ xuất từ Google Docs ----------------
FLOAT_RE = re.compile(r"^-?\d+\.\d+$")
for el in doc.element.iter():
    if not isinstance(el.tag, str):
        continue
    for name, val in list(el.attrib.items()):
        if name.startswith("{%s}" % W) and FLOAT_RE.match(val):
            el.set(name, str(int(round(float(val)))))
for pgmar in doc.element.iter(q("pgMar")):
    if pgmar.get(q("gutter")) is None:
        pgmar.set(q("gutter"), "0")

doc.save(OUT)
print("Đã lưu", OUT)
print("Số tài liệu tham khảo:", len(order))
for k in order:
    print(f"  [{NUM[k]}] {k}" + (f" (cũ [{OLD_NUM[k]}])" if k in OLD_NUM else ""))
